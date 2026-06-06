"""
AI智能文档润色/翻译助手 v2.2
Word进Word出 · PDF智能摘要 · 现代UI
"""

import io
import json
import logging
import os
import re
import time
from typing import Generator, Optional, List, Dict

import requests
import streamlit as st

# ============================================================
# 配置
# ============================================================

API_PROVIDERS = {
    "通义千问": {
        "base_url": "https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions",
        "models": ["qwen-turbo", "qwen-plus", "qwen-max"],
        "default_model": "qwen-turbo",
    },
    "DeepSeek": {
        "base_url": "https://api.deepseek.com/chat/completions",
        "models": ["deepseek-chat", "deepseek-reasoner"],
        "default_model": "deepseek-chat",
    },
}

ROLE_TEMPLATES = {
    "严谨学术风": {
        "icon": "🎓",
        "color": "#3b82f6",
        "desc": "学术论文润色",
        "prompt": "你是一位资深学术编辑。对以下文本进行深度润色：提升学术规范性，优化句子结构，修正语法错误，使用恰当学术词汇。直接返回润色后的文本，不要添加任何解释说明。",
    },
    "硅谷极客风": {
        "icon": "⚡",
        "color": "#06b6d4",
        "desc": "技术文档优化",
        "prompt": "你是一位技术文档专家。优化表达使其简洁专业，统一技术术语，删除冗余。直接返回优化后的文本，不要添加任何解释说明。",
    },
    "简洁商务风": {
        "icon": "💼",
        "color": "#8b5cf6",
        "desc": "商务报告改写",
        "prompt": "你是一位商务沟通顾问。改写为简洁专业的商务表达，突出关键信息。直接返回改写后的文本，不要添加任何解释说明。",
    },
    "中英翻译": {
        "icon": "🌏",
        "color": "#10b981",
        "desc": "学术翻译",
        "prompt": "你是一位学术翻译专家。进行中英或英中专业翻译，保持术语一致性。直接返回翻译结果，不要添加任何解释说明。",
    },
    "代码解释": {
        "icon": "💻",
        "color": "#f59e0b",
        "desc": "代码注释",
        "prompt": "你是一位技术导师。为代码添加清晰详尽的中文注释。直接返回带注释的代码，不要添加任何解释说明。",
    },
}

# PDF摘要专用配置
PDF_CONFIG = {
    "chunk_size": 3000,       # 每块最大字符数
    "chunk_overlap": 200,     # 块间重叠字符数（保证上下文连贯）
    "max_chunks": 10,         # 最大处理块数（防止超长PDF）
}

APP_CONFIG = {
    "max_input_length": 8000,
    "default_temperature": 0.3,
    "timeout_seconds": 120,
}

# ============================================================
# API 客户端
# ============================================================

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger(__name__)


class LLMClient:
    def __init__(self, provider: str, api_key: str, model: Optional[str] = None):
        if provider not in API_PROVIDERS:
            raise ValueError(f"不支持的API提供商: {provider}")
        self.provider = provider
        self.api_key = api_key
        self.config = API_PROVIDERS[provider]
        self.model = model or self.config["default_model"]
        self.base_url = self.config["base_url"]

    def process_sync(self, text: str, system_prompt: str) -> str:
        headers = {"Authorization": f"Bearer {self.api_key}", "Content-Type": "application/json"}
        payload = {
            "model": self.model,
            "messages": [{"role": "system", "content": system_prompt}, {"role": "user", "content": text}],
            "temperature": APP_CONFIG["default_temperature"],
            "stream": False,
        }
        response = requests.post(self.base_url, headers=headers, json=payload,
                                 timeout=APP_CONFIG["timeout_seconds"])
        if response.status_code != 200:
            raise ValueError(f"API请求失败 [{response.status_code}]")
        return response.json()["choices"][0]["message"]["content"]

    def stream_process(self, text: str, system_prompt: str) -> Generator[str, None, None]:
        headers = {"Authorization": f"Bearer {self.api_key}", "Content-Type": "application/json"}
        payload = {
            "model": self.model,
            "messages": [{"role": "system", "content": system_prompt}, {"role": "user", "content": text}],
            "temperature": APP_CONFIG["default_temperature"],
            "stream": True,
        }
        try:
            response = requests.post(self.base_url, headers=headers, json=payload,
                                     timeout=APP_CONFIG["timeout_seconds"], stream=True)
        except requests.exceptions.Timeout:
            raise TimeoutError("请求超时")
        except requests.exceptions.ConnectionError:
            raise ConnectionError("网络连接失败")

        if response.status_code != 200:
            raise ValueError(f"API请求失败 [{response.status_code}]")

        for line in response.iter_lines(decode_unicode=True):
            if not line or not line.startswith("data: "):
                continue
            data = line[6:]
            if data == "[DONE]":
                break
            try:
                chunk = json.loads(data)
                content = chunk.get("choices", [{}])[0].get("delta", {}).get("content", "")
                if content:
                    yield content
            except:
                continue


# ============================================================
# Word 文档处理
# ============================================================

def read_docx_structured(file_bytes: bytes) -> List[Dict]:
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = []
    for p in doc.paragraphs:
        if p.text.strip():
            paragraphs.append({
                "text": p.text.strip(),
                "style": p.style.name if p.style else "Normal"
            })
    return paragraphs


def merge_short_paragraphs(paragraphs: List[Dict], min_chars: int = 100, max_chars: int = 1500) -> List[Dict]:
    """合并短段落，减少API调用次数，提升处理速度"""
    if not paragraphs:
        return []
    merged = []
    buffer_text = ""
    buffer_style = "Normal"
    for para in paragraphs:
        text = para["text"]
        style = para["style"]
        # 标题段落单独保留
        if "Heading" in style or "标题" in style:
            if buffer_text:
                merged.append({"text": buffer_text.strip(), "style": buffer_style})
                buffer_text = ""
            merged.append(para)
            continue
        # 累积正文段落
        if len(buffer_text) + len(text) < max_chars:
            buffer_text += "\n\n" + text if buffer_text else text
            buffer_style = style
        else:
            if buffer_text:
                merged.append({"text": buffer_text.strip(), "style": buffer_style})
            buffer_text = text
            buffer_style = style
    if buffer_text:
        merged.append({"text": buffer_text.strip(), "style": buffer_style})
    return merged


def save_docx_structured(paragraphs: List[Dict], title: str = "润色结果") -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Microsoft YaHei'
    font.size = Pt(11)

    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x1e, 0x3a, 0x5f)

    doc.add_paragraph("")
    from datetime import datetime
    info = doc.add_paragraph(f"处理时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.runs[0].font.size = Pt(9)
    info.runs[0].font.color.rgb = RGBColor(0x94, 0xa3, 0xb8)
    doc.add_paragraph("")
    doc.add_paragraph("─" * 40)
    doc.add_paragraph("")

    for para in paragraphs:
        text = para.get("text", "").strip()
        style_name = para.get("style", "Normal")
        if not text:
            continue
        if "Heading" in style_name or "标题" in style_name:
            level = 1 if "1" in style_name else 2 if "2" in style_name else 3 if "3" in style_name else 1
            doc.add_heading(text, level=level)
        else:
            doc.add_paragraph(text)

    doc.add_paragraph("")
    doc.add_paragraph("─" * 40)
    footer = doc.add_paragraph("由 AI智能文档润色助手 生成")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor(0x94, 0xa3, 0xb8)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output.getvalue()


# ============================================================
# PDF 处理（核心功能）
# ============================================================

def read_pdf(file_bytes: bytes) -> str:
    import PyPDF2
    reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
    texts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            texts.append(text.strip())
    return "\n\n".join(texts)


def split_text_to_chunks(text: str, chunk_size: int = 3000, overlap: int = 200) -> List[str]:
    """将长文本分割成有重叠的块，保证上下文连贯"""
    if len(text) <= chunk_size:
        return [text]

    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        if end >= len(text):
            chunks.append(text[start:])
            break
        # 在句子边界处切割
        cut_pos = text.rfind("。", start, end)
        if cut_pos == -1:
            cut_pos = text.rfind("\n", start, end)
        if cut_pos == -1:
            cut_pos = end
        chunks.append(text[start:cut_pos + 1])
        start = cut_pos + 1 - overlap  # 重叠保证上下文

    return chunks[:PDF_CONFIG["max_chunks"]]


def summarize_chunk(client: LLMClient, chunk: str, chunk_index: int, total: int) -> str:
    """对单块内容生成摘要"""
    prompt = f"""你是一位学术摘要专家。请对以下文档片段（第{chunk_index}/{total}部分）进行摘要：

要求：
1. 提取核心观点、关键数据、重要结论
2. 用中文输出，分点列出
3. 简明扼要，每点不超过30字
4. 如果片段包含研究方法，请简要说明

文档片段：
{chunk}"""
    return client.process_sync(chunk, prompt)


def merge_summaries(summaries: List[str]) -> str:
    """合并多个块摘要为最终摘要"""
    merged = "\n\n".join(summaries)
    return merged


def generate_final_summary(client: LLMClient, merged_summary: str) -> str:
    """从合并摘要中提炼最终摘要"""
    prompt = f"""你是一位学术摘要专家。请根据以下分块摘要，生成一份完整的文档摘要报告：

要求：
1. 包含：【文档主题】→【核心观点】→【研究方法】→【关键结论】→【主要贡献】
2. 每个部分用中文输出，清晰分点
3. 总字数控制在300-500字
4. 语言严谨、专业

分块摘要：
{merged_summary}"""
    return client.process_sync(merged_summary, prompt)


# ============================================================
# 工具函数
# ============================================================

def count_stats(text: str) -> dict:
    clean = text.strip()
    cn = len(re.findall(r"[\u4e00-\u9fff]", clean))
    return {"总字符": len(clean), "中文": cn, "段落": max(len([p for p in clean.split("\n\n") if p.strip()]), 1)}


# ============================================================
# Streamlit 界面
# ============================================================

st.set_page_config(page_title="AI智能文档润色助手", page_icon="📝", layout="wide")

st.markdown("""
<style>
.block-container { padding: 1rem 2rem !important; max-width: 100% !important; }
.hero-section {
    background: linear-gradient(135deg, #0f172a 0%, #1e3a5f 50%, #0d9488 100%);
    border-radius: 16px; padding: 2rem 2.5rem; margin-bottom: 1.5rem;
    color: white; box-shadow: 0 8px 32px rgba(15,23,42,0.3);
}
.hero-title { font-size: 2rem !important; font-weight: 700; margin: 0; }
.hero-subtitle { font-size: 0.95rem; color: #94a3b8; margin-top: 0.3rem; }
.role-card {
    border: 2px solid #e2e8f0; border-radius: 14px; padding: 1rem 0.8rem;
    text-align: center; background: white; transition: all 0.3s;
}
.role-card:hover { transform: translateY(-3px); box-shadow: 0 8px 24px rgba(0,0,0,0.08); }
.role-card.active { border-color: #3b82f6; background: #eff6ff; }
.role-icon { font-size: 1.8rem; margin-bottom: 0.3rem; }
.role-name { font-weight: 600; font-size: 0.85rem; color: #1e293b; }
.role-desc { font-size: 0.72rem; color: #94a3b8; margin-top: 2px; }
.pdf-card {
    background: linear-gradient(135deg, #fef2f2, #fee2e2);
    border: 2px solid #ef4444; border-radius: 14px; padding: 1.2rem;
    text-align: center; cursor: pointer; transition: all 0.3s;
}
.pdf-card:hover { transform: translateY(-3px); box-shadow: 0 8px 24px rgba(239,68,68,0.15); }
.stat-box {
    background: linear-gradient(135deg, #f0f9ff, #e0f2fe);
    border-radius: 10px; padding: 8px 12px; text-align: center;
}
.stat-num { font-size: 1.3rem; font-weight: 700; color: #0369a1; }
.stat-label { font-size: 0.7rem; color: #64748b; }
.result-area {
    background: #f8fafc; border-left: 4px solid #0d9488;
    border-radius: 0 10px 10px 0; padding: 1.2rem;
    min-height: 380px; max-height: 600px; overflow-y: auto;
    font-size: 0.95rem; line-height: 1.8; color: #334155;
}
#MainMenu, footer, header, .stDeployButton { display: none !important; }
section[data-testid="stSidebar"] { min-width: 280px !important; }
</style>
""", unsafe_allow_html=True)

# ============================================================
# 会话状态
# ============================================================

if "result" not in st.session_state:
    st.session_state.result = ""
if "structured_paras" not in st.session_state:
    st.session_state.structured_paras = []
if "is_word_mode" not in st.session_state:
    st.session_state.is_word_mode = False
if "is_pdf_mode" not in st.session_state:
    st.session_state.is_pdf_mode = False
if "count" not in st.session_state:
    st.session_state.count = 0
if "role" not in st.session_state:
    st.session_state.role = "严谨学术风"
if "stop_flag" not in st.session_state:
    st.session_state.stop_flag = False
if "processing" not in st.session_state:
    st.session_state.processing = False

api_key = ""

# ============================================================
# 侧边栏
# ============================================================

with st.sidebar:
    st.header("⚙️ 设置")
    provider = st.selectbox("API提供商", list(API_PROVIDERS.keys()), index=0)

    default_key = ""
    if "DASHSCOPE_API_KEY" in st.secrets:
        default_key = st.secrets["DASHSCOPE_API_KEY"]
    elif "DASHSCOPE_API_KEY" in os.environ:
        default_key = os.environ["DASHSCOPE_API_KEY"]

    if default_key:
        api_key = default_key
        st.success("✅ API Key已加载")
    else:
        api_key = st.text_input("API Key", type="password", placeholder="粘贴你的Key")

    model = st.selectbox("模型", API_PROVIDERS[provider]["models"], index=0)
    st.markdown("---")
    st.caption("💡 推荐通义千问，免费额度充足")

# ============================================================
# Hero 区域
# ============================================================

st.markdown("""
<div class="hero-section">
    <div style="display:flex;justify-content:space-between;align-items:center;">
        <div>
            <div class="hero-title">📝 AI 智能文档润色助手</div>
            <div class="hero-subtitle">Word进Word出 · PDF智能摘要 · 保留原文格式</div>
        </div>
        <div style="font-size:2.5rem;font-weight:700;opacity:0.9;">v2.2</div>
    </div>
</div>
""", unsafe_allow_html=True)

if not api_key:
    st.warning("⚠️ 请在左侧边栏输入API Key")
    st.stop()

# ============================================================
# 功能选择区（润色角色 + PDF摘要）
# ============================================================

st.subheader("🎭 选择功能")

# 第一行：润色角色
cols = st.columns(len(ROLE_TEMPLATES))
for col, (key, info) in zip(cols, ROLE_TEMPLATES.items()):
    with col:
        active = "active" if st.session_state.role == key and not st.session_state.is_pdf_mode else ""
        st.markdown(f"""
        <div class="role-card {active}">
            <div class="role-icon">{info['icon']}</div>
            <div class="role-name">{key}</div>
            <div class="role-desc">{info['desc']}</div>
        </div>""", unsafe_allow_html=True)
        if st.button("选择", key=f"sel_{key}", use_container_width=True):
            st.session_state.role = key
            st.session_state.is_word_mode = False
            st.session_state.is_pdf_mode = False
            st.session_state.structured_paras = []
            st.session_state.result = ""
            st.rerun()

# 第二行：PDF摘要（独立入口）
st.markdown("<div style='margin-top:0.8rem'></div>", unsafe_allow_html=True)
pdf_cols = st.columns([1, 3, 1])
with pdf_cols[1]:
    pdf_active = "style='border-color:#ef4444;background:#fef2f2;'" if st.session_state.is_pdf_mode else ""
    st.markdown(f"""
    <div class="pdf-card" {pdf_active}>
        <div style="font-size:2rem;margin-bottom:0.3rem">📄</div>
        <div style="font-weight:700;font-size:1.1rem;color:#991b1b">PDF 智能摘要</div>
        <div style="font-size:0.8rem;color:#b91c1c;margin-top:0.3rem">上传PDF文档 → 自动分段 → 生成结构化摘要报告</div>
    </div>
    """, unsafe_allow_html=True)
    if st.button("📄 进入PDF摘要模式", key="pdf_mode_btn", use_container_width=True):
        st.session_state.is_pdf_mode = True
        st.session_state.is_word_mode = False
        st.session_state.role = ""
        st.session_state.structured_paras = []
        st.session_state.result = ""
        st.rerun()

# 当前模式显示
if st.session_state.is_pdf_mode:
    st.info("📄 **PDF智能摘要模式** — 上传PDF文档，AI自动分段生成结构化摘要报告")
else:
    info = ROLE_TEMPLATES[st.session_state.role]
    st.info(f"{info['icon']} **{st.session_state.role}** — {info['desc']}")

st.markdown("---")

# ============================================================
# 输入区域
# ============================================================

st.subheader("📥 输入")

input_text = ""
source_paras = []

if st.session_state.is_pdf_mode:
    # PDF摘要模式
    file = st.file_uploader("📄 上传PDF文档", type=["pdf"], label_visibility="collapsed")
    if file:
        try:
            with st.spinner("🔍 正在读取PDF内容..."):
                input_text = read_pdf(file.getvalue())
            stats = count_stats(input_text)
            st.success(f"✅ PDF读取成功: {file.name} · {stats['总字符']}字符 · {stats['段落']}段落")

            # 显示PDF内容预览（可折叠）
            with st.expander("📋 查看PDF提取内容"):
                st.text(input_text[:2000] + ("..." if len(input_text) > 2000 else ""))
        except Exception as e:
            st.error(f"❌ PDF读取失败: {e}")

else:
    # 润色模式
    method = st.radio("方式", ["直接输入", "上传文件 (txt/docx)"], horizontal=True, label_visibility="collapsed")

    if method.startswith("直接输入"):
        input_text = st.text_area("文本", placeholder="粘贴需要处理的文本...", height=280,
                                   max_chars=APP_CONFIG["max_input_length"], label_visibility="collapsed")
    else:
        file = st.file_uploader("上传", type=["txt", "docx"], label_visibility="collapsed")
        if file:
            try:
                if file.name.endswith(".docx"):
                    source_paras = read_docx_structured(file.getvalue())
                    input_text = "\n\n".join([p["text"] for p in source_paras])
                    st.session_state.is_word_mode = True
                    st.success(f"✅ Word: {file.name} · {len(source_paras)}个段落")
                else:
                    input_text = file.getvalue().decode("utf-8")
                    st.success(f"✅ 文本: {file.name}")
            except Exception as e:
                st.error(f"❌ 读取失败: {e}")

if input_text.strip() and not st.session_state.is_pdf_mode:
    stats = count_stats(input_text)
    cols = st.columns(4)
    for c, (k, v) in zip(cols, stats.items()):
        c.markdown(f"<div class='stat-box'><div class='stat-num'>{v}</div><div class='stat-label'>{k}</div></div>", unsafe_allow_html=True)

# 开始/停止按钮
col_btn1, col_btn2 = st.columns([1, 1])
with col_btn1:
    process = st.button("🚀 开始处理", type="primary", use_container_width=True,
                        disabled=not (api_key and input_text.strip()) or st.session_state.processing)
with col_btn2:
    if st.session_state.processing:
        if st.button("⏹️ 停止处理", type="secondary", use_container_width=True):
            st.session_state.stop_flag = True
            st.session_state.processing = False
            st.rerun()

st.markdown("---")

# ============================================================
# 处理逻辑
# ============================================================

st.subheader("📤 处理结果")

result_placeholder = st.empty()

if process and api_key and input_text.strip():
    try:
        client = LLMClient(provider=provider, api_key=api_key, model=model)

        # ========== PDF摘要模式 ==========
        if st.session_state.is_pdf_mode:
            st.session_state.stop_flag = False
            st.session_state.processing = True

            # 1. 分段
            chunks = split_text_to_chunks(input_text)
            st.info(f"📑 文档已分割为 {len(chunks)} 个片段进行处理")
            progress = st.progress(0)
            status = st.empty()

            # 2. 逐块摘要
            chunk_summaries = []
            for i, chunk in enumerate(chunks):
                if st.session_state.stop_flag:
                    status.warning("⏹️ 处理已停止（部分完成）")
                    break

                status.info(f"🤖 正在处理第 {i+1}/{len(chunks)} 个片段...")
                try:
                    summary = summarize_chunk(client, chunk, i+1, len(chunks))
                    chunk_summaries.append(summary)
                except Exception as e:
                    chunk_summaries.append(f"[第{i+1}片段处理失败: {str(e)}]")
                progress.progress((i + 1) / len(chunks))

            # 3. 合并并生成最终摘要
            if chunk_summaries and not st.session_state.stop_flag:
                status.info("📝 正在生成最终摘要报告...")
                merged = merge_summaries(chunk_summaries)
                final_summary = generate_final_summary(client, merged)
                st.session_state.result = final_summary
            elif chunk_summaries:
                st.session_state.result = merge_summaries(chunk_summaries)
            else:
                st.session_state.result = ""

            st.session_state.count += 1
            st.session_state.processing = False
            progress.empty()
            status.empty()

            if st.session_state.stop_flag:
                st.warning(f"⏹️ 已停止，完成 {len(chunk_summaries)}/{len(chunks)} 个片段")
            else:
                st.success(f"✅ PDF摘要完成！处理了 {len(chunks)} 个片段")

        # ========== Word润色模式 ==========
        elif st.session_state.is_word_mode and source_paras:
            # 合并短段落减少API调用
            merged_paras = merge_short_paragraphs(source_paras)
            st.info(f"📊 原文 {len(source_paras)} 个段落 → 合并为 {len(merged_paras)} 个处理块（提升速度）")

            # 初始化停止标志
            st.session_state.stop_flag = False
            st.session_state.processing = True

            progress = st.progress(0)
            status = st.empty()
            prompt = ROLE_TEMPLATES[st.session_state.role]["prompt"]

            polished_paras = []
            for i, para in enumerate(merged_paras):
                if st.session_state.stop_flag:
                    status.warning("⏹️ 处理已停止（部分完成）")
                    break

                status.info(f"🤖 正在处理第 {i+1}/{len(merged_paras)} 块...")

                if para["text"].strip():
                    try:
                        polished = client.process_sync(para["text"], prompt)
                        polished_paras.append({"text": polished.strip(), "style": para["style"]})
                    except Exception as e:
                        st.warning(f"第{i+1}块处理失败: {e}")
                        polished_paras.append(para)
                else:
                    polished_paras.append(para)
                progress.progress((i + 1) / len(merged_paras))

            st.session_state.structured_paras = polished_paras
            st.session_state.result = "\n\n".join([p["text"] for p in polished_paras])
            st.session_state.count += 1
            st.session_state.processing = False
            progress.empty()
            status.empty()

            if st.session_state.stop_flag:
                st.warning(f"⏹️ 已停止，完成 {len(polished_paras)}/{len(merged_paras)} 块")
            else:
                st.success(f"✅ Word润色完成！处理了 {len(merged_paras)} 个块")

        # ========== 普通模式 ==========
        else:
            prompt = ROLE_TEMPLATES[st.session_state.role]["prompt"]
            result_placeholder.info("🤖 AI处理中...")
            full = ""
            t0 = time.time()
            for chunk in client.stream_process(input_text, prompt):
                full += chunk
                safe = full.replace("<", "&lt;").replace(">", "&gt;")
                result_placeholder.markdown(f"<div class='result-area'>{safe}</div>", unsafe_allow_html=True)
            st.session_state.result = full
            st.session_state.structured_paras = []
            st.session_state.count += 1
            st.success(f"✅ 完成！耗时 {time.time()-t0:.1f}秒")

    except Exception as e:
        result_placeholder.error(f"❌ 失败: {e}")

elif st.session_state.result:
    safe = st.session_state.result.replace("<", "&lt;").replace(">", "&gt;")
    result_placeholder.markdown(f"<div class='result-area'>{safe}</div>", unsafe_allow_html=True)
else:
    result_placeholder.info("👆 输入后点击「开始处理」")

# ============================================================
# 导出区域
# ============================================================

if st.session_state.result:
    st.markdown("---")
    st.subheader("💾 导出结果")

    # Word模式：导出完整格式Word
    if st.session_state.is_word_mode and st.session_state.structured_paras:
        try:
            docx_data = save_docx_structured(st.session_state.structured_paras,
                                              title=f"{st.session_state.role}结果")
            st.download_button("📥 下载完整Word文档", docx_data,
                               file_name=f"润色结果_{st.session_state.role}.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               use_container_width=True)
        except Exception as e:
            st.error(f"Word导出失败: {e}")

    # PDF模式：导出Word格式摘要报告
    elif st.session_state.is_pdf_mode:
        c1, c2 = st.columns(2)
        with c1:
            st.text_area("📋 复制文本", st.session_state.result, height=80, label_visibility="collapsed")
            st.caption("👆 Ctrl+A 全选 → Ctrl+C 复制")
        with c2:
            try:
                # 生成Word格式摘要报告
                paras = [{"text": p, "style": "Normal"} for p in st.session_state.result.split("\n\n") if p.strip()]
                docx_data = save_docx_structured(paras, title="PDF文档摘要报告")
                st.download_button("📄 导出Word摘要报告", docx_data,
                                   file_name="PDF摘要报告.docx",
                                   mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                   use_container_width=True)
            except:
                txt_data = st.session_state.result.encode("utf-8")
                st.download_button("📝 导出TXT", txt_data, "PDF摘要报告.txt", "text/plain", use_container_width=True)

    # 普通模式：文本导出
    else:
        c1, c2 = st.columns(2)
        with c1:
            st.text_area("📋 复制文本", st.session_state.result, height=80, label_visibility="collapsed")
            st.caption("👆 Ctrl+A 全选 → Ctrl+C 复制")
        with c2:
            txt_data = st.session_state.result.encode("utf-8")
            st.download_button("📝 导出TXT", txt_data,
                               file_name=f"结果_{st.session_state.role}.txt",
                               mime="text/plain", use_container_width=True)

st.markdown("""
<div style="text-align:center;color:#94a3b8;font-size:0.8rem;margin-top:2rem;padding:1rem;border-top:1px solid #e2e8f0;">
    🎓 东北大学秦皇岛分校 · 计算机与通信工程学院 · 启航科创项目 | v2.2
</div>
""", unsafe_allow_html=True)
